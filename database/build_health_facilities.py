#!/usr/bin/env python3
"""Prépare les lignes du classeur pour le catalogue PostgreSQL ientier."""

from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


CATEGORY = {
    "CS": "Centre de santé",
    "CCS": "Centre communautaire de santé",
    "HCR": "Hôpital communautaire de référence",
    "HD": "Hôpital départemental",
    "HU": "Hôpital universitaire",
    "Autres hôpitaux": "Autres hôpitaux",
    "PPS": "Point de prestation de services",
    "Point Fixe": "Point fixe",
}
OWNERSHIP = {"Publique", "Privé à but lucratif", "Privé sans but lucratif", "Mixte"}
REVIEW_TOKENS = {
    "CS", "CCS", "CSL", "CDS", "HCR", "HD", "HU", "SSPE", "PPS",
    "FOSREF", "DASH", "OFATMA", "GHESKIO", "MSF", "SADA", "HHF",
    "FONDEFH", "MEBSH", "SOSAKA", "RHEMA", "UNDH", "ODECOM",
    "AFSC", "NDPS", "DECADE", "ESMI", "FONDHAPE-S", "IDADEE",
    "COMEFASA", "CSIN", "HELMT", "CEMAD", "CMSMO", "CEDIMED",
    "CEMEFCO", "HNDP", "HBSBB", "COHASA", "HCAH", "POZ",
    "COSECO", "FOSADEC", "FACEDHA", "FHED", "OPORHA", "HUEH",
    "SODEC", "APROSIFA", "UMO", "CEMEAH", "CDSP", "OMEC",
    "CESMED", "FHHH", "CHOSCAL", "CHAPI", "CEMEPID", "CPFO",
    "MUSAC", "SSFM", "NF", "PROFAMIL", "FONFRED", "MIC",
    "SOS", "ND", "JMC", "AME", "LCM", "LIB", "APV", "HAH",
    "OMES", "OBCG", "ASCOSADEC", "ASAD", "IBS", "CD", "SA",
    "HEB", "FAME", "PEREO", "SAA",
}
TOKEN_PATTERN = re.compile(r"(?<!\w)(?:" + "|".join(
    re.escape(token) for token in sorted(REVIEW_TOKENS, key=len, reverse=True)
) + r")(?!\w)")
HEADERS = (
    "ID", "Département", "Arrondissement", "Commune ", "Section Communale",
    "Nom de l'institution", "Situation de fonctionnement", "Catégorie", "Statut",
    "Code Institution",
)
FIELDS = (
    "source_sha256", "source_row", "legacy_id", "source_code", "source_name",
    "display_name", "department", "arrondissement", "commune", "communal_section",
    "category", "ownership_status", "operation_status", "source_operation_status",
    "name_review_required", "record_review_required", "published",
)


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def sql_literal(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def load_rows(source: Path) -> tuple[list[dict], dict]:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Le classeur doit contenir une seule feuille de données.")
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        if tuple(next(values)[:10]) != HEADERS:
            raise ValueError("Les en-têtes du classeur ont changé.")
        rows = []
        for row_number, values_row in enumerate(values, start=2):
            if not any(value is not None for value in values_row):
                continue
            source_row = dict(zip(HEADERS, values_row[:10]))
            name = clean(source_row["Nom de l'institution"])
            department = clean(source_row["Département"])
            arrondissement = clean(source_row["Arrondissement"])
            commune = clean(source_row["Commune "])
            if not all((name, department, arrondissement, commune)):
                raise ValueError(f"Ligne {row_number} sans nom ou localisation obligatoire.")
            raw_category = clean(source_row["Catégorie"])
            if raw_category is not None and raw_category not in CATEGORY:
                raise ValueError(f"Catégorie inconnue à la ligne {row_number}: {raw_category}")
            raw_ownership = clean(source_row["Statut"])
            if raw_ownership is not None and raw_ownership not in OWNERSHIP:
                raise ValueError(f"Statut juridique inconnu à la ligne {row_number}: {raw_ownership}")
            raw_operation = clean(source_row["Situation de fonctionnement"])
            operation = {
                "Fonctionnelle": "Fonctionnelle",
                "Non Fonctionnelle": "Non fonctionnelle",
            }.get(raw_operation, "À vérifier")
            code = clean(source_row["Code Institution"])
            legacy_id = source_row["ID"]
            if legacy_id is not None and not isinstance(legacy_id, int):
                raise ValueError(f"ID non entier à la ligne {row_number}")
            rows.append({
                "source_sha256": digest,
                "source_row": row_number,
                "legacy_id": legacy_id,
                "source_code": code,
                "source_name": name,
                "display_name": name,
                "department": department,
                "arrondissement": arrondissement,
                "commune": commune,
                "communal_section": clean(source_row["Section Communale"]),
                "category": CATEGORY.get(raw_category),
                "ownership_status": raw_ownership,
                "operation_status": operation,
                "source_operation_status": raw_operation,
                "name_review_required": bool(TOKEN_PATTERN.search(name)),
                "record_review_required": False,
                "published": False,
            })
    finally:
        workbook.close()

    code_counts = collections.Counter(row["source_code"] for row in rows if row["source_code"])
    anomalies = []
    for row in rows:
        issues = []
        if not row["source_code"]:
            issues.append("code manquant")
        elif code_counts[row["source_code"]] > 1:
            issues.append("code partagé")
        if row["category"] is None:
            issues.append("catégorie manquante")
        if row["ownership_status"] is None:
            issues.append("statut juridique manquant")
        if row["communal_section"] is None:
            issues.append("section communale manquante")
        if row["operation_status"] == "À vérifier":
            issues.append("fonctionnement à vérifier")
        row["record_review_required"] = bool(issues)
        if issues:
            anomalies.append({
                "source_row": row["source_row"], "name": row["source_name"],
                "issues": issues,
            })
    report = {
        "source_file": source.name,
        "source_sha256": digest,
        "row_count": len(rows),
        "department_counts": dict(sorted(collections.Counter(
            row["department"] for row in rows
        ).items())),
        "category_counts": dict(sorted(collections.Counter(
            row["category"] or "Non renseignée" for row in rows
        ).items())),
        "operation_counts": dict(sorted(collections.Counter(
            row["operation_status"] for row in rows
        ).items())),
        "missing_code_count": sum(not row["source_code"] for row in rows),
        "shared_codes": dict(sorted((code, count) for code, count in code_counts.items()
                                    if count > 1)),
        "name_review_count": sum(row["name_review_required"] for row in rows),
        "name_reviews": [
            {
                "source_row": row["source_row"],
                "source_code": row["source_code"],
                "source_name": row["source_name"],
                "sigles_to_review": TOKEN_PATTERN.findall(row["source_name"]),
            }
            for row in rows if row["name_review_required"]
        ],
        "record_review_count": len(anomalies),
        "anomalies": anomalies,
    }
    return rows, report


def write_sql(rows: list[dict], destination: Path) -> None:
    with destination.open("w", encoding="utf-8") as output:
        output.write("-- Données extraites du classeur ; exécuter après health_facilities_schema.sql.\n")
        output.write("-- Chaque ligne est identifiée par l'empreinte du fichier et son numéro Excel.\n")
        output.write("BEGIN;\n")
        for offset in range(0, len(rows), 100):
            batch = rows[offset:offset + 100]
            output.write("INSERT INTO ientier.health_facilities (\n  ")
            output.write(", ".join(FIELDS) + "\n) VALUES\n")
            output.write(",\n".join("  (" + ", ".join(sql_literal(row[field]) for field in FIELDS) + ")"
                                     for row in batch))
            output.write("\nON CONFLICT (source_sha256, source_row) DO NOTHING;\n")
        output.write("COMMIT;\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Classeur Excel d'origine")
    args = parser.parse_args()
    rows, report = load_rows(args.source)
    if len(rows) != 1072:
        raise ValueError(f"1072 lignes attendues ; {len(rows)} trouvées")
    output_dir = Path(__file__).resolve().parent
    with (output_dir / "health_facilities.csv").open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    write_sql(rows, output_dir / "health_facilities_seed.sql")
    (output_dir / "health_facilities_audit.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"{report['row_count']} lignes préparées ; {report['name_review_count']} noms à revoir ; "
          f"{report['record_review_count']} fiches à vérifier")


if __name__ == "__main__":
    main()
